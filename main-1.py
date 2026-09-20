import os
import json
import logging
import asyncio
from aiohttp import web

from aiogram import Bot, Dispatcher, F
from aiogram.filters import CommandStart
from aiogram.types import Message, FSInputFile
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.fsm.storage.memory import MemoryStorage

from pptx import Presentation
from pptx.util import Inches, Pt

import google.generativeai as genai

# ---------- SOZLAMALAR ----------
BOT_TOKEN = os.environ.get("BOT_TOKEN")
GEMINI_API_KEY = os.environ.get("GEMINI_API_KEY")

if not BOT_TOKEN:
    raise RuntimeError("BOT_TOKEN environment o'zgaruvchisi topilmadi!")
if not GEMINI_API_KEY:
    raise RuntimeError("GEMINI_API_KEY environment o'zgaruvchisi topilmadi!")

logging.basicConfig(level=logging.INFO)

bot = Bot(token=BOT_TOKEN)
dp = Dispatcher(storage=MemoryStorage())

genai.configure(api_key=GEMINI_API_KEY)
gemini_model = genai.GenerativeModel("gemini-2.0-flash")


# ---------- HOLATLAR (FSM) ----------
class SlideForm(StatesGroup):
    waiting_topic = State()


# ---------- AI ORQALI SLAYD MATNI GENERATSIYA QILISH ----------
def generate_slides_content(topic: str, slide_count: int = 8) -> list:
    """
    AI dan mavzu bo'yicha slaydlar matnini JSON formatida so'raydi.
    Har bir slayd: {"title": "...", "bullets": ["...", "..."]}
    """
    prompt = f"""Mavzu: "{topic}"

Shu mavzu bo'yicha {slide_count} ta slaydlik prezentatsiya tarkibini tayyorla.
Birinchi slayd — sarlavha slaydi (faqat title, bullets bo'sh massiv).
Qolganlari — har birida sarlavha va 3-5 ta qisqa punkt (bullet) bo'lsin.
Matn o'zbek tilida bo'lsin.

FAQAT quyidagi JSON formatida javob ber, boshqa hech qanday matn, izoh yoki
markdown belgilarisiz (```json kabi belgilar ham bo'lmasin):

[
  {{"title": "Slayd sarlavhasi", "bullets": ["punkt 1", "punkt 2"]}},
  ...
]
"""

    response = gemini_model.generate_content(prompt)

    raw_text = response.text.strip()

    # Ehtiyot chorasi: agar model baribir ```json bilan o'rasa, tozalaymiz
    raw_text = raw_text.replace("```json", "").replace("```", "").strip()

    slides_data = json.loads(raw_text)
    return slides_data


# ---------- PPTX FAYL YASASH ----------
def build_presentation(topic: str, slides_data: list, output_path: str):
    prs = Presentation()

    # 1-slayd: sarlavha
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    slide.shapes.title.text = slides_data[0].get("title", topic)
    if len(slide.placeholders) > 1:
        slide.placeholders[1].text = topic

    # Qolgan slaydlar: sarlavha + bullet list
    bullet_layout = prs.slide_layouts[1]
    for item in slides_data[1:]:
        slide = prs.slides.add_slide(bullet_layout)
        slide.shapes.title.text = item.get("title", "")

        body = slide.placeholders[1]
        tf = body.text_frame
        tf.clear()

        bullets = item.get("bullets", [])
        for i, bullet_text in enumerate(bullets):
            if i == 0:
                p = tf.paragraphs[0]
            else:
                p = tf.add_paragraph()
            p.text = bullet_text
            p.font.size = Pt(20)

    prs.save(output_path)


# ---------- HANDLERLAR ----------
@dp.message(CommandStart())
async def cmd_start(message: Message, state: FSMContext):
    await state.set_state(SlideForm.waiting_topic)
    await message.answer(
        "Salom! Men prezentatsiya (.pptx) tayyorlab beraman.\n\n"
        "Menga mavzuni yozing, masalan:\n"
        "«Sun'iy intellekt kelajakda ta'limga ta'siri»"
    )


@dp.message(SlideForm.waiting_topic, F.text)
async def handle_topic(message: Message, state: FSMContext):
    topic = message.text.strip()
    processing_msg = await message.answer("⏳ Prezentatsiya tayyorlanmoqda, biroz kuting...")

    try:
        slides_data = generate_slides_content(topic, slide_count=8)

        output_path = f"/tmp/{message.from_user.id}_presentation.pptx"
        build_presentation(topic, slides_data, output_path)

        file = FSInputFile(output_path, filename="Prezentatsiya.pptx")
        await message.answer_document(file, caption=f"✅ Tayyor! Mavzu: {topic}")

        os.remove(output_path)

    except Exception as e:
        logging.exception("Xatolik yuz berdi")
        await message.answer(
            "❌ Xatolik yuz berdi, qaytadan urinib ko'ring.\n"
            f"(texnik tafsilot: {e})"
        )
    finally:
        await processing_msg.delete()


# ---------- ISHGA TUSHIRISH ----------
async def handle_ping(request):
    return web.Response(text="Bot ishlayapti")


async def start_web_server():
    """
    Render.com 'Web Service' turi doim ochiq port kutadi.
    Shuning uchun botning yonida mayda http server ham ishga tushiramiz.
    """
    app = web.Application()
    app.router.add_get("/", handle_ping)

    port = int(os.environ.get("PORT", 8080))

    runner = web.AppRunner(app)
    await runner.setup()
    site = web.TCPSite(runner, host="0.0.0.0", port=port)
    await site.start()
    logging.info(f"Web server {port}-portda ishga tushdi")


async def main():
    await start_web_server()
    await dp.start_polling(bot)


if __name__ == "__main__":
    asyncio.run(main())
